"""Secure, bounded resume file validation and text extraction."""

from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import pymupdf

from domain.schemas import DocumentChunk, DocumentSection, ParsedDocument


MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_PDF_PAGES = 50
MAX_DOCX_UNCOMPRESSED_SIZE = 25 * 1024 * 1024
DEFAULT_CHUNK_SIZE = 1_000
DEFAULT_CHUNK_OVERLAP = 150
PARSER_VERSION = "1"


class ResumeParseError(ValueError):
    pass


def validate_resume_file(filename, data):
    """Validate type, signature, and size without trusting browser metadata."""
    suffix = Path(filename).suffix.lower()
    if suffix not in {".pdf", ".docx"}:
        raise ResumeParseError("Only PDF and DOCX files are supported.")
    if not data:
        raise ResumeParseError("The uploaded file is empty.")
    if len(data) > MAX_FILE_SIZE:
        raise ResumeParseError("The file exceeds the 10 MB size limit.")
    if suffix == ".pdf" and not data.startswith(b"%PDF-"):
        raise ResumeParseError("The file extension is PDF, but its signature is invalid.")
    if suffix == ".docx" and not data.startswith(b"PK"):
        raise ResumeParseError("The file extension is DOCX, but its signature is invalid.")
    return suffix


def extract_text_from_bytes(filename, data):
    """Return plain text for legacy callers.

    New code should use :func:`parse_document_from_bytes` for source provenance.
    """
    suffix = validate_resume_file(filename, data)
    return _extract_pdf(data) if suffix == ".pdf" else _extract_docx(data)


def parse_document_from_bytes(filename, data, *, chunk_size=DEFAULT_CHUNK_SIZE,
                              chunk_overlap=DEFAULT_CHUNK_OVERLAP):
    """Parse a supported file into source sections and retrieval chunks.

    Chunking is section-aware: no chunk spans a PDF page or a DOCX heading/table
    region. DOCX has no reliable pagination in ``python-docx``, so its sections
    intentionally have ``page_number=None``.
    """
    suffix = validate_resume_file(filename, data)
    _validate_chunking_options(chunk_size, chunk_overlap)
    if suffix == ".pdf":
        parts = _extract_pdf_sections(data)
        media_type = "application/pdf"
        parser_name = "pymupdf"
    else:
        parts = _extract_docx_sections(data)
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        parser_name = "python-docx"
    return _build_parsed_document(
        filename, media_type, parser_name, parts, chunk_size, chunk_overlap
    )


def _extract_pdf(data):
    try:
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            if doc.needs_pass:
                raise ResumeParseError("Password-protected PDFs are not supported.")
            if doc.page_count > MAX_PDF_PAGES:
                raise ResumeParseError(f"PDF exceeds the {MAX_PDF_PAGES}-page limit.")
            return "\n".join(page.get_text("text") for page in doc).strip()
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError("The PDF is corrupted or could not be read.") from exc


def _extract_pdf_sections(data):
    try:
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            if doc.needs_pass:
                raise ResumeParseError("Password-protected PDFs are not supported.")
            if doc.page_count > MAX_PDF_PAGES:
                raise ResumeParseError(f"PDF exceeds the {MAX_PDF_PAGES}-page limit.")
            sections = []
            for number, page in enumerate(doc, start=1):
                page_text = page.get_text("text").strip()
                if page_text:
                    sections.append({
                        "kind": "page", "title": f"Page {number}",
                        "text": page_text, "page_number": number,
                    })
            return sections
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError("The PDF is corrupted or could not be read.") from exc


def _extract_docx(data):
    try:
        from docx import Document
    except ImportError as exc:
        raise ResumeParseError("DOCX support requires the python-docx package.") from exc
    try:
        with ZipFile(BytesIO(data)) as archive:
            if "word/document.xml" not in archive.namelist():
                raise ResumeParseError("The uploaded file is not a valid DOCX document.")
            if sum(item.file_size for item in archive.infolist()) > MAX_DOCX_UNCOMPRESSED_SIZE:
                raise ResumeParseError("The DOCX expands beyond the safe processing limit.")
        document = Document(BytesIO(data))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            paragraphs.extend(cell.text for row in table.rows for cell in row.cells)
        return "\n".join(paragraphs).strip()
    except ResumeParseError:
        raise
    except (BadZipFile, OSError, ValueError, KeyError) as exc:
        raise ResumeParseError("The DOCX file is corrupted or could not be read.") from exc


def _extract_docx_sections(data):
    try:
        from docx import Document
    except ImportError as exc:
        raise ResumeParseError("DOCX support requires the python-docx package.") from exc
    try:
        with ZipFile(BytesIO(data)) as archive:
            if "word/document.xml" not in archive.namelist():
                raise ResumeParseError("The uploaded file is not a valid DOCX document.")
            if sum(item.file_size for item in archive.infolist()) > MAX_DOCX_UNCOMPRESSED_SIZE:
                raise ResumeParseError("The DOCX expands beyond the safe processing limit.")
        document = Document(BytesIO(data))
        sections = []
        current_title = None
        body_lines = []

        def flush_body():
            if body_lines:
                sections.append({
                    "kind": "heading" if current_title else "body",
                    "title": current_title,
                    "text": "\n".join(body_lines).strip(),
                    "page_number": None,
                })
                body_lines.clear()

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = getattr(paragraph.style, "name", "") or ""
            if style_name.casefold().startswith("heading"):
                flush_body()
                sections.append({
                    "kind": "heading", "title": text, "text": text,
                    "page_number": None,
                })
                current_title = text
            else:
                body_lines.append(text)
        flush_body()
        for table in document.tables:
            cells = [cell.text.strip() for row in table.rows for cell in row.cells if cell.text.strip()]
            if cells:
                sections.append({
                    "kind": "table", "title": current_title,
                    "text": "\n".join(cells), "page_number": None,
                })
        return sections
    except ResumeParseError:
        raise
    except (BadZipFile, OSError, ValueError, KeyError) as exc:
        raise ResumeParseError("The DOCX file is corrupted or could not be read.") from exc


def extract_text_from_pdf(pdf_file):
    """
    Extract text from an uploaded PDF file.
    """
    return extract_text_from_bytes(getattr(pdf_file, "name", "resume.pdf"), pdf_file.read())


def _validate_chunking_options(chunk_size, chunk_overlap):
    if not isinstance(chunk_size, int) or chunk_size < 1:
        raise ValueError("chunk_size must be a positive integer")
    if not isinstance(chunk_overlap, int) or chunk_overlap < 0 or chunk_overlap >= chunk_size:
        raise ValueError("chunk_overlap must be non-negative and smaller than chunk_size")


def _build_parsed_document(filename, media_type, parser_name, parts, chunk_size, chunk_overlap):
    text_parts, sections, chunks = [], [], []
    source_offset = 0
    for number, part in enumerate(parts, start=1):
        if text_parts:
            text_parts.append("\n")
            source_offset += 1
        section_text = part["text"]
        section_id = f"section-{number:04d}"
        section = DocumentSection(
            section_id=section_id, kind=part["kind"], title=part["title"],
            text=section_text, source_start=source_offset,
            source_end=source_offset + len(section_text), page_number=part["page_number"],
        )
        sections.append(section)
        for chunk_number, (start, end) in enumerate(
            _split_section(section_text, chunk_size, chunk_overlap), start=1
        ):
            chunk_text = section_text[start:end]
            chunks.append(DocumentChunk(
                chunk_id=f"{section_id}-chunk-{chunk_number:04d}", section_id=section_id,
                text=chunk_text, source_start=source_offset + start,
                source_end=source_offset + end, page_number=part["page_number"],
            ))
        text_parts.append(section_text)
        source_offset += len(section_text)
    return ParsedDocument(
        filename=filename, media_type=media_type, parser_name=parser_name,
        parser_version=PARSER_VERSION, text="".join(text_parts), sections=sections, chunks=chunks,
    )


def _split_section(text, chunk_size, chunk_overlap):
    """Deterministic fallback splitter preferring whitespace over mid-word cuts."""
    if not text:
        return []
    spans, start, text_length = [], 0, len(text)
    while start < text_length:
        end = min(start + chunk_size, text_length)
        if end < text_length:
            boundary = text.rfind(" ", start + 1, end + 1)
            if boundary > start:
                end = boundary
        if end <= start:
            end = min(start + chunk_size, text_length)
        spans.append((start, end))
        if end == text_length:
            break
        start = max(end - chunk_overlap, start + 1)
    return spans
