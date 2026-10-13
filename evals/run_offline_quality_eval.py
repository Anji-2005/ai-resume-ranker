"""Versioned synthetic parser, privacy, score, and schema evaluation."""

import argparse
from io import BytesIO
import json
from pathlib import Path
from time import perf_counter

from docx import Document
import pymupdf

from utils.genai_analyzer import GenAIResponseError, validate_enrichment
from utils.pdf_parser import parse_document_from_bytes
from utils.privacy import anonymize_resume_text
from utils.scoring import calculate_score_breakdown


def _document_bytes(case):
    if case["format"] == "docx":
        output = BytesIO()
        document = Document()
        document.add_heading(case.get("heading", "Section"), level=1)
        document.add_paragraph(case["text"])
        document.save(output)
        return f"{case['id']}.docx", output.getvalue()
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), case["text"])
    content = document.tobytes()
    document.close()
    return f"{case['id']}.pdf", content


def _structured_payload(case):
    return {
        "candidate_summary": "The fictional document states Python API work.",
        "improvement_suggestions": ["Add measurable outcomes if accurate."],
        "interview_questions": [
            {"question": case["question"], "reason": case["reason"]},
            {"question": "What tradeoffs did you consider?", "reason": "Explore technical decisions."},
            {"question": "How was the service tested?", "reason": "Verify engineering practice."},
        ],
        "requirement_evidence": [{
            "requirement": "Python", "status": "supported", "evidence": "Python API",
            "explanation": "Exact evidence is present.",
        }],
    }


def evaluate(dataset_path):
    dataset = json.loads(dataset_path.read_text())
    started = perf_counter()

    parser_passes = 0
    for case in dataset["parser_cases"]:
        filename, content = _document_bytes(case)
        parsed = parse_document_from_bytes(filename, content)
        exact_offsets = all(parsed.text[chunk.source_start:chunk.source_end] == chunk.text for chunk in parsed.chunks)
        expected_page = all(chunk.page_number == case["expected_page"] for chunk in parsed.chunks)
        parser_passes += bool(case["text"] in parsed.text and exact_offsets and expected_page)

    redacted = total_sensitive = 0
    for case in dataset["redaction_cases"]:
        output = anonymize_resume_text(case["text"])
        for value in case["sensitive_values"]:
            total_sensitive += 1
            redacted += value not in output

    score_passes = 0
    for case in dataset["score_cases"]:
        result = calculate_score_breakdown(
            case["semantic_score"], case["resume_text"], case["required_skills"],
            case["preferred_skills"], case["candidate_years"], case["minimum_years"],
        )
        score_passes += result["final_score"] == case["expected_final_score"]

    schema_passes = sensitive_accepted = sensitive_total = 0
    for case in dataset["structured_output_cases"]:
        try:
            validate_enrichment(_structured_payload(case), "Python API", ["Python"])
            outcome = "accept"
        except GenAIResponseError:
            outcome = "reject"
        schema_passes += outcome == case["expected"]
        if case["sensitive"]:
            sensitive_total += 1
            sensitive_accepted += outcome == "accept"

    return {
        "dataset_version": dataset["dataset_version"],
        "method": "deterministic_offline_quality",
        "parser_fidelity": round(parser_passes / len(dataset["parser_cases"]), 4),
        "redaction_recall": round(redacted / total_sensitive, 4),
        "score_parity": round(score_passes / len(dataset["score_cases"]), 4),
        "structured_output_accuracy": round(schema_passes / len(dataset["structured_output_cases"]), 4),
        "sensitive_content_acceptance_rate": round(sensitive_accepted / sensitive_total, 4),
        "latency_ms": round((perf_counter() - started) * 1_000, 3),
        "live_model_calls": 0,
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--dataset", type=Path, default=Path("data/offline_quality_v1.json"))
    parser.add_argument("--output", type=Path, default=Path("evals/latest_offline_quality.json"))
    args = parser.parse_args()
    result = evaluate(args.dataset)
    args.output.write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps(result, indent=2))
