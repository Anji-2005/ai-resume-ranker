"""Reproducible PostgreSQL ingestion-worker benchmark using fictional documents."""

import argparse
from datetime import UTC, datetime
import hashlib
from io import BytesIO
import json
from pathlib import Path
from time import perf_counter

from docx import Document
import pymupdf
from sqlalchemy import delete
from sqlalchemy.orm import sessionmaker
from sqlalchemy import create_engine

from db.models import IngestionJob, StoredDocument
from ingestion.worker import IngestionWorker
from repositories.ingestion_jobs import IngestionJobRepository
from services.embeddings import SentenceTransformerEmbedder


TENANT = "synthetic-ingestion-benchmark"


class DeterministicBenchmarkEmbedder:
    model_name = "deterministic-benchmark-embedder"
    model_version = "1"

    @staticmethod
    def embed(texts):
        vectors = []
        for text in texts:
            digest = hashlib.sha256(text.encode()).digest()
            vectors.append([((digest[index % len(digest)] / 255.0) - 0.5) for index in range(384)])
        return vectors


def _document_bytes(case):
    if case["format"] == "docx":
        output = BytesIO()
        document = Document()
        document.add_heading("Experience", level=1)
        for paragraph in case["paragraphs"]:
            document.add_paragraph(paragraph)
        document.save(output)
        return "document.docx", output.getvalue()
    document = pymupdf.open()
    page = document.new_page()
    page.insert_textbox(pymupdf.Rect(40, 40, 550, 800), "\n\n".join(case["paragraphs"]), fontsize=11)
    data = document.tobytes()
    document.close()
    return "document.pdf", data


def _percentile(values, ratio):
    ordered = sorted(values)
    if not ordered:
        return 0.0
    return ordered[max(0, int(len(ordered) * ratio + 0.999999) - 1)]


def run(database_url, dataset_path, *, real_embedder=False):
    dataset = json.loads(dataset_path.read_text())
    engine = create_engine(database_url, pool_pre_ping=True)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    embedder_factory = SentenceTransformerEmbedder if real_embedder else DeterministicBenchmarkEmbedder
    total_bytes = 0
    try:
        with factory.begin() as session:
            session.execute(delete(IngestionJob).where(IngestionJob.tenant_id == TENANT))
            session.execute(delete(StoredDocument).where(StoredDocument.tenant_id == TENANT))
            jobs = IngestionJobRepository(session)
            model = embedder_factory()
            for case in dataset["documents"]:
                filename, data = _document_bytes(case)
                total_bytes += len(data)
                content_hash = hashlib.sha256(data).hexdigest()
                jobs.enqueue(
                    tenant_id=TENANT, idempotency_key=f"benchmark-{dataset['dataset_version']}-{case['id']}",
                    request_hash=content_hash, created_by="token-sha256:synthetic-benchmark",
                    job_type="ingest", source_key=f"benchmark-{case['id']}", filename=filename,
                    content_hash=content_hash, payload=data, embedding_model=model.model_name,
                    embedding_model_version=model.model_version, max_attempts=1,
                )
        worker = IngestionWorker(factory, embedder_factory=embedder_factory, lease_seconds=600)
        started = perf_counter()
        processed = 0
        while worker.run_once():
            processed += 1
        elapsed = perf_counter() - started
        with factory() as session:
            rows = list(session.query(IngestionJob).filter(IngestionJob.tenant_id == TENANT).all())
            latencies = [row.processing_latency_ms for row in rows if row.processing_latency_ms is not None]
            completed = sum(row.status == "completed" for row in rows)
            failed = sum(row.status == "failed" for row in rows)
        return {
            "dataset_version": dataset["dataset_version"],
            "benchmark_version": "postgres-worker-ingestion-v1",
            "generated_at": datetime.now(UTC).isoformat(),
            "embedding_mode": "real MiniLM" if real_embedder else "deterministic fake (parser/DB/worker only)",
            "documents": len(rows),
            "processed_claims": processed,
            "completed": completed,
            "failed": failed,
            "success_rate": completed / len(rows) if rows else 0.0,
            "total_input_bytes": total_bytes,
            "wall_seconds": elapsed,
            "throughput_documents_per_second": completed / elapsed if elapsed else 0.0,
            "throughput_bytes_per_second": total_bytes / elapsed if elapsed else 0.0,
            "mean_processing_latency_ms": sum(latencies) / len(latencies) if latencies else None,
            "p50_processing_latency_ms": _percentile(latencies, 0.50),
            "p95_processing_latency_ms": _percentile(latencies, 0.95),
            "live_model_calls": 0,
            "limitations": [
                "Single worker and six small synthetic documents; this is not a concurrency or load test.",
                ("MiniLM embedding runtime is included." if real_embedder else
                 "A deterministic fake embedder excludes production model load and inference latency."),
                "Local PostgreSQL and filesystem/cache conditions materially affect timings.",
            ],
        }
    finally:
        with factory.begin() as session:
            session.execute(delete(IngestionJob).where(IngestionJob.tenant_id == TENANT))
            session.execute(delete(StoredDocument).where(StoredDocument.tenant_id == TENANT))
        engine.dispose()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--database-url", required=True)
    parser.add_argument("--dataset", type=Path, default=Path("data/synthetic_ingestion_v1.json"))
    parser.add_argument("--output", type=Path, default=Path("evals/latest_ingestion_benchmark.json"))
    parser.add_argument("--real-embedder", action="store_true")
    args = parser.parse_args()
    result = run(args.database_url, args.dataset, real_embedder=args.real_embedder)
    args.output.write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps({key: value for key, value in result.items() if key not in {"limitations"}}, indent=2))


if __name__ == "__main__":
    main()
