from io import BytesIO
import unittest

import pymupdf
import pandas as pd
from docx import Document

from utils.candidate_analyzer import estimate_years_experience
from utils.pdf_parser import (
    ResumeParseError, extract_text_from_bytes, parse_document_from_bytes,
    validate_resume_file,
)
from utils.privacy import (
    anonymize_resume_text,
    safe_candidate_label,
    safe_failure_record,
    sanitize_audit_records,
)
from utils.scoring import calculate_score_breakdown, normalize_weights, review_recommendation
from utils.skills_extractor import (
    extract_skills, find_skill_evidence, find_skill_evidence_with_provenance,
    parse_skill_list,
)


class SkillExtractionTests(unittest.TestCase):
    def test_aliases_are_canonicalized(self):
        self.assertEqual(parse_skill_list("python; powerbi, sklearn"),
                         ["Python", "Power BI", "scikit-learn"])

    def test_word_boundaries_avoid_nosql_false_positive(self):
        self.assertNotIn("SQL", extract_skills("Built a NoSQL document store"))

    def test_evidence_is_returned_from_source_text(self):
        evidence = find_skill_evidence("Built APIs with Python. Used SQL daily.", ["Python", "SQL"])
        self.assertIn("Python", evidence["Python"][0])
        self.assertIn("SQL", evidence["SQL"][0])


class ScoringTests(unittest.TestCase):
    def test_weights_are_normalized(self):
        weights = normalize_weights({"semantic": 40, "required_skills": 40,
                                     "preferred_skills": 10, "experience": 10})
        self.assertAlmostEqual(sum(weights.values()), 1.0)

    def test_unconfigured_optional_criteria_have_zero_weight(self):
        result = calculate_score_breakdown(
            50, "Python developer", ["Python", "SQL"], [], 0, 0,
            {"semantic": 40, "required_skills": 35, "preferred_skills": 10, "experience": 15},
        )
        self.assertEqual(result["weights"]["preferred_skills"], 0)
        self.assertEqual(result["weights"]["experience"], 0)
        self.assertAlmostEqual(result["final_score"], 50.0)

    def test_review_label_is_not_hiring_decision(self):
        self.assertEqual(review_recommendation(85, 80), "Priority review")
        self.assertNotIn(review_recommendation(40, 20), {"Hire", "Reject"})


class ParsingAndPrivacyTests(unittest.TestCase):
    def test_pdf_signature_is_checked(self):
        with self.assertRaises(ResumeParseError):
            validate_resume_file("resume.pdf", b"not a PDF")

    def test_text_pdf_is_extracted(self):
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 72), "Python developer with five years of experience")
        data = document.tobytes()
        document.close()
        self.assertIn("Python developer", extract_text_from_bytes("resume.pdf", data))

    def test_docx_is_extracted(self):
        output = BytesIO()
        document = Document()
        document.add_paragraph("SQL analyst with 3 years of experience")
        document.save(output)
        self.assertIn("SQL analyst", extract_text_from_bytes("resume.docx", output.getvalue()))

    def test_pdf_parse_has_page_sections_chunks_and_exact_offsets(self):
        document = pymupdf.open()
        first = document.new_page()
        first.insert_text((72, 72), "Python analysis on page one")
        second = document.new_page()
        second.insert_text((72, 72), "SQL reporting on page two")
        data = document.tobytes()
        document.close()

        parsed = parse_document_from_bytes("resume.pdf", data, chunk_size=12, chunk_overlap=3)

        self.assertEqual(parsed.media_type, "application/pdf")
        self.assertEqual([section.page_number for section in parsed.sections], [1, 2])
        self.assertTrue(all(chunk.page_number in {1, 2} for chunk in parsed.chunks))
        self.assertGreater(len(parsed.chunks), len(parsed.sections))
        for section in parsed.sections:
            self.assertEqual(parsed.text[section.source_start:section.source_end], section.text)
        for chunk in parsed.chunks:
            self.assertEqual(parsed.text[chunk.source_start:chunk.source_end], chunk.text)

    def test_docx_parse_uses_heading_sections_without_invented_pages(self):
        output = BytesIO()
        document = Document()
        document.add_heading("Experience", level=1)
        document.add_paragraph("Built Python services and SQL reports.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Skill"
        table.cell(0, 1).text = "Python"
        document.save(output)

        parsed = parse_document_from_bytes("resume.docx", output.getvalue(), chunk_size=500)

        self.assertIn("Experience", parsed.text)
        self.assertIn("Built Python services", parsed.text)
        self.assertIn("Skill\nPython", parsed.text)
        self.assertTrue(any(section.kind == "heading" for section in parsed.sections))
        self.assertTrue(any(section.kind == "table" for section in parsed.sections))
        self.assertTrue(all(section.page_number is None for section in parsed.sections))
        self.assertTrue(all(chunk.page_number is None for chunk in parsed.chunks))

    def test_chunking_is_section_aware_and_configurable(self):
        document = pymupdf.open()
        for text in ("A" * 50, "B" * 50):
            page = document.new_page()
            page.insert_text((72, 72), text)
        data = document.tobytes()
        document.close()

        parsed = parse_document_from_bytes("resume.pdf", data, chunk_size=20, chunk_overlap=5)
        section_ids = {chunk.section_id for chunk in parsed.chunks}
        self.assertEqual(section_ids, {"section-0001", "section-0002"})
        self.assertTrue(all(len(chunk.text) <= 20 for chunk in parsed.chunks))
        with self.assertRaises(ValueError):
            parse_document_from_bytes("resume.pdf", data, chunk_size=10, chunk_overlap=10)

    def test_structured_skill_evidence_has_exact_provenance(self):
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 72), "Built Python APIs. Used SQL for reporting.")
        data = document.tobytes()
        document.close()
        parsed = parse_document_from_bytes("resume.pdf", data, chunk_size=500)

        evidence = find_skill_evidence_with_provenance(parsed, ["Python", "SQL"])

        self.assertEqual(evidence["Python"][0].page_number, 1)
        self.assertEqual(evidence["SQL"][0].section_id, "section-0001")
        for items in evidence.values():
            for item in items:
                self.assertEqual(parsed.text[item.source_start:item.source_end], item.excerpt)

    def test_excel_engine_is_available(self):
        output = BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            pd.DataFrame([{"candidate": "Candidate 01", "score": 80}]).to_excel(
                writer, index=False
            )
        self.assertTrue(output.getvalue().startswith(b"PK"))

    def test_direct_identifiers_are_redacted(self):
        redacted = anonymize_resume_text("Jane Doe\njane@example.com\n+91 98765 43210")
        self.assertNotIn("Jane Doe", redacted)
        self.assertNotIn("jane@example.com", redacted)
        self.assertNotIn("98765", redacted)
        self.assertEqual(safe_candidate_label("jane.pdf", 2, True), "Candidate 02")

    def test_blind_review_records_do_not_expose_source_filenames(self):
        filename = "jane-doe-private-resume.pdf"
        failure = safe_failure_record(
            filename, 1, f"Could not parse {filename}", blind_review=True
        )
        results, failures = sanitize_audit_records(
            [{"candidate": "Candidate 01", "source_filename": filename}],
            [{"file": filename, "error": f"Could not parse {filename}"}],
            blind_review=True,
        )

        self.assertEqual(failure["file"], "Candidate 01")
        self.assertNotIn(filename, str(failure))
        self.assertEqual(results[0]["source_filename"], "Candidate 01")
        self.assertEqual(failures[0]["file"], "Candidate 01")
        self.assertNotIn(filename, str((results, failures)))

    def test_non_blind_records_preserve_source_filenames(self):
        filename = "fictional-resume.pdf"
        failure = safe_failure_record(filename, 1, "parse failed", blind_review=False)
        results, failures = sanitize_audit_records(
            [{"candidate": filename, "source_filename": filename}],
            [failure],
            blind_review=False,
        )

        self.assertEqual(results[0]["source_filename"], filename)
        self.assertEqual(failures[0]["file"], filename)

    def test_experience_requires_explicit_claim(self):
        self.assertEqual(estimate_years_experience("5 years of experience in analytics"), 5)
        self.assertEqual(estimate_years_experience("Graduated in 2020"), 0)


if __name__ == "__main__":
    unittest.main()
