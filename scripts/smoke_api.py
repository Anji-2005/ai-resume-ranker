"""Synthetic HTTP smoke test for the running Compose API.

Uses only fictional documents and never calls a live LLM: it expects grounded
generation to be unconfigured (503) or policy-denied (403) unless
``SMOKE_EXPECT_GENERATION=enabled`` is set explicitly.
"""

from io import BytesIO
import os
import time
from uuid import uuid4

from docx import Document
import httpx


def synthetic_docx(*paragraphs):
    output = BytesIO()
    document = Document()
    document.add_heading("Experience", level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(output)
    return output.getvalue()


DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def wait_for_job(client, job_id, headers, timeout_seconds=120):
    deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < deadline:
        response = client.get(f"/v1/ingestion-jobs/{job_id}", headers=headers)
        assert response.status_code == 200, response.text
        job = response.json()
        if job["status"] in {"completed", "failed"}:
            return job
        time.sleep(1)
    raise AssertionError(f"ingestion job did not reach a terminal state within {timeout_seconds}s")


def main():
    base_url = os.getenv("SMOKE_API_URL", "http://localhost:8000")
    tenant = os.getenv("SMOKE_TENANT_ID", "synthetic-tenant")
    token = os.getenv("SMOKE_BEARER_TOKEN", "local-validation-token")
    reader_token = os.getenv("SMOKE_READER_TOKEN")
    headers = {"X-Tenant-ID": tenant, "Authorization": f"Bearer {token}"}
    checks = []

    with httpx.Client(base_url=base_url, timeout=120) as client:
        health = client.get("/v1/health")
        assert health.status_code == 200 and health.json()["database"] == "ok", health.text
        checks.append("health")

        unauthorized = client.post("/v1/documents/search", json={"query": "Python"}, headers={"X-Tenant-ID": tenant})
        assert unauthorized.status_code == 401, unauthorized.text
        forbidden = client.post("/v1/documents/search", json={"query": "Python"},
                                headers={"X-Tenant-ID": "other-tenant", "Authorization": f"Bearer {token}"})
        assert forbidden.status_code == 403, forbidden.text
        checks.append("auth+tenant denial")

        session = client.get("/v1/session", headers=headers)
        assert session.status_code == 200, session.text
        generation = session.json()["generation"]
        checks.append(f"session(generation={generation})")

        data = synthetic_docx("Built Python FastAPI services with PostgreSQL and Docker for a fictional project.",
                              "Contact: smoke.candidate@example.com")
        ingest = client.post("/v1/documents", files={"file": ("Fictional Person.docx", data, DOCX)}, headers=headers)
        assert ingest.status_code == 201, ingest.text
        document_id = ingest.json()["document_id"]
        repeat = client.post("/v1/documents", files={"file": ("again.docx", data, DOCX)}, headers=headers)
        assert repeat.status_code == 201 and repeat.json()["created"] is False, repeat.text
        checks.append("ingest+idempotency")

        background_data = synthetic_docx(
            "Maintained Python data pipelines for a fictional warehouse.",
        )
        background_headers = {**headers, "Idempotency-Key": "synthetic-smoke-" + uuid4().hex}
        queued = client.post(
            "/v1/ingestion-jobs", headers=background_headers,
            files={"file": ("Fictional Background Person.docx", background_data, DOCX)},
        )
        assert queued.status_code in {200, 202}, queued.text
        replay = client.post(
            "/v1/ingestion-jobs", headers=background_headers,
            files={"file": ("Different Identity.docx", background_data, DOCX)},
        )
        assert replay.status_code == 200 and replay.json()["job_id"] == queued.json()["job_id"], replay.text
        hidden = client.get(
            f"/v1/ingestion-jobs/{queued.json()['job_id']}",
            headers={"X-Tenant-ID": "other-tenant", "Authorization": f"Bearer {token}"},
        )
        assert hidden.status_code in {403, 404}, hidden.text
        background = wait_for_job(client, queued.json()["job_id"], headers)
        assert background["status"] == "completed", background
        background_document_id = background["document_id"]
        metrics = client.get("/v1/ingestion-jobs/metrics", headers=headers)
        assert metrics.status_code == 200 and metrics.json()["counts"]["completed"] >= 1, metrics.text
        checks.append("background ingest+idempotency+status+tenant isolation")

        listing = client.get("/v1/documents", headers=headers)
        assert listing.status_code == 200 and any(d["document_id"] == document_id for d in listing.json()["documents"])
        assert "Fictional Person" not in listing.text
        detail = client.get(f"/v1/documents/{document_id}", headers=headers)
        assert detail.status_code == 200 and detail.json()["chunks"], detail.text
        assert "smoke.candidate@example.com" not in detail.text
        checks.append("list+detail(redacted)")

        search = client.post("/v1/documents/search", json={"query": "Python FastAPI", "document_ids": [document_id]},
                             headers=headers)
        assert search.status_code == 200 and search.json(), search.text
        assert all(item["document_id"] == document_id for item in search.json())
        legacy = client.get("/v1/documents/search", params={"query": "Python FastAPI"}, headers=headers)
        assert legacy.status_code == 200, legacy.text
        checks.append("hybrid search")

        expected = {"unconfigured": 503, "policy_denied": 403}.get(generation)
        if os.getenv("SMOKE_EXPECT_GENERATION") != "enabled":
            assert expected is not None, "refusing to run a live provider call during smoke validation"
            qa = client.post("/v1/documents/qa", json={"question": "What was built?", "document_ids": [document_id]},
                             headers=headers)
            assert qa.status_code == expected, qa.text
            workflow = client.post("/v1/workflows/recruiter-brief", headers=headers, json={
                "description": "Python engineer", "document_ids": [document_id]})
            assert workflow.status_code == expected, workflow.text
            investigation = client.post(
                "/v1/investigations",
                headers={**headers, "Idempotency-Key": "synthetic-smoke-investigation"},
                json={
                    "objective": "Investigate Python API evidence",
                    "document_ids": [document_id],
                    "requirements": ["Python"],
                },
            )
            assert investigation.status_code == expected, investigation.text
            checks.append(f"generation gate fails closed ({expected})")

        history = client.get("/v1/workflows", headers=headers)
        assert history.status_code == 200, history.text
        investigation_history = client.get("/v1/investigations", headers=headers)
        assert investigation_history.status_code == 200, investigation_history.text
        checks.append("workflow+investigation history")

        if reader_token:
            reader = {"X-Tenant-ID": tenant, "Authorization": f"Bearer {reader_token}"}
            denied = client.delete(f"/v1/documents/{document_id}", headers=reader)
            assert denied.status_code == 403, denied.text
            checks.append("reader cannot delete")

        unauthorized_delete = client.delete(f"/v1/documents/{document_id}", headers={"X-Tenant-ID": tenant})
        assert unauthorized_delete.status_code == 401, unauthorized_delete.text
        deleted = client.delete(f"/v1/documents/{document_id}", headers=headers)
        assert deleted.status_code == 204, deleted.text
        after = client.post("/v1/documents/search", json={"query": "Python FastAPI"}, headers=headers)
        assert all(item["document_id"] != document_id for item in after.json())
        assert client.get(f"/v1/documents/{document_id}", headers=headers).status_code == 404
        background_deleted = client.delete(f"/v1/documents/{background_document_id}", headers=headers)
        assert background_deleted.status_code == 204, background_deleted.text
        checks.append("physical deletion")

    print("API smoke passed: " + ", ".join(checks))


if __name__ == "__main__":
    main()
